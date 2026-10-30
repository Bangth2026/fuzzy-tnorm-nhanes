from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = "Times New Roman"
thin = Side(style="thin", color="999999")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="D9E2F3")
FILL = PatternFill("solid", fgColor="FFF2CC")   # ô chuyên gia điền
SUG = PatternFill("solid", fgColor="EDEDED")    # gợi ý tham khảo
wrap = Alignment(wrap_text=True, vertical="top")

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        x = ws.cell(row, c); x.font = Font(name=F, bold=True, size=11); x.fill = HEAD; x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); x.border = B

wb = Workbook()

# ---------------- Sheet 1: Hướng dẫn ----------------
ws = wb.active; ws.title = "Huong dan"
lines = [
 ("PHIẾU DUYỆT HÀM THÀNH VIÊN CÁC KHÁI NIỆM", True, 14),
 ("Bài báo: Keep the Grades, Choose the Product: Auditing Structure and Multi-hop t-norm Inference in a Population Fuzzy Knowledge Graph of Depressive Symptoms (dự kiến gửi Journal of Biomedical Informatics)", False, 11),
 ("Người duyệt: Phan Thị Vân An (chuyên gia tâm lý lâm sàng)   ·   Người gửi: Trương Hải Bằng", False, 11),
 ("", False, 11),
 ("Mục đích", True, 12),
 ("Đồ thị tri thức mờ gồm 24 khái niệm trích từ NHANES 2007–2018. Mỗi khái niệm có một hàm thành viên μ ∈ [0, 1] chuyển câu trả lời khảo sát thành mức độ. 9 triệu chứng PHQ-9 dùng thang chuẩn; 15 yếu tố nguy cơ dùng ngưỡng do tác giả đặt và cần chuyên gia lâm sàng duyệt trước khi chạy lại toàn bộ pipeline.", False, 11),
 ("", False, 11),
 ("Cách điền (sheet 'Phieu duyet')", True, 12),
 ("1. Đọc cột F (hàm hiện tại, đúng như mã nguồn clean2.py đang chạy).", False, 11),
 ("2. Cột G–H (nền xám) là GỢI Ý THAM KHẢO do trợ lý soạn, kèm nguồn. Đây không phải ý kiến của chuyên gia; cô có thể bỏ qua hoàn toàn.", False, 11),
 ("3. Chỉ điền các ô nền vàng: cột I chọn 'Giữ' hoặc 'Sửa' (danh sách thả xuống); nếu 'Sửa', ghi hàm/ngưỡng mới ở cột J theo dạng ở ví dụ dưới; cột K ghi nguồn cô dựa vào; cột L ghi chú tự do.", False, 11),
 ("4. Sheet 'Cau hoi lam sang': trả lời ngắn 4 câu hỏi về diễn giải kết quả (vài câu mỗi ý là đủ).", False, 11),
 ("5. Ký tên, ghi ngày ở cuối sheet 'Cau hoi lam sang' và gửi lại file.", False, 11),
 ("", False, 11),
 ("Ví dụ một dòng đã điền (minh họa định dạng, không phải đề xuất)", True, 12),
 ("Khái niệm: short_sleep  |  Quyết định: Sửa  |  Hàm mới: 0 khi ngủ ≥ 7 giờ, 1 khi ngủ ≤ 6 giờ, tuyến tính ở giữa  |  Nguồn: Watson et al., Sleep 2015  |  Ghi chú: ngưỡng 5 giờ quá chặt cho mẫu dân số chung", False, 11),
 ("", False, 11),
 ("Cách viết hàm mới", True, 12),
 ("• Hàm dốc tuyến tính: '0 tại a, 1 tại b' (ví dụ BMI: 0 tại 25, 1 tại 30).", False, 11),
 ("• Hàm bậc thang theo mã trả lời: liệt kê mã → giá trị (ví dụ: hút hằng ngày = 1; thỉnh thoảng = 0,5; không = 0).", False, 11),
 ("• Nếu cần tách theo giới hoặc tuổi, ghi rõ từng nhóm.", False, 11),
 ("", False, 11),
 ("Sau khi nhận phiếu", True, 12),
 ("Tác giả sửa clean2.py theo phiếu và chạy lại toàn bộ pipeline theo README (clean2 → analysis2 main/nulls/rq3 → power2 → nresp_check → stats2). Mọi con số trong bài lấy từ results2/paper_stats.json; không sửa số bằng tay. Phiếu đã điền được lưu làm hồ sơ minh chứng cho mục Methods và CRediT.", False, 11),
]
for i, (t, b, s) in enumerate(lines, 1):
    c = ws.cell(i, 1, t); c.font = Font(name=F, bold=b, size=s); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120
for i in (2, 6, 10, 11, 16, 24):
    ws.row_dimensions[i].height = 48

# ---------------- Sheet 2: Phiếu duyệt ----------------
ws = wb.create_sheet("Phieu duyet")
hdr = ["STT", "Nhóm", "Khái niệm (mã)", "Tên tiếng Việt", "Biến NHANES", "Hàm hiện tại (μ = 0 → μ = 1)",
       "Gợi ý tham khảo (trợ lý soạn, không phải ý kiến chuyên gia)", "Nguồn của gợi ý",
       "Quyết định (Giữ / Sửa)", "Hàm / ngưỡng mới (nếu Sửa)", "Nguồn chuyên gia dựa vào", "Ghi chú"]
ws.append(hdr); style_header(ws, 1, len(hdr))
R = [
 ("Triệu chứng", "9 mục PHQ-9 (anhedonia … suicidal_ideation)", "9 triệu chứng PHQ-9", "DPQ010–DPQ090",
  "Điểm mục 0–3 chia cho 3 (0; 0,33; 0,67; 1). Không đặt ngưỡng.",
  "Giữ: đây là thang chuẩn của công cụ. Lưu ý mục 9 (ý nghĩ tự hại) có tỉ lệ dương thấp nên các cạnh nối với nó kém ổn định hơn.",
  "Kroenke et al. 2001 [R1]"),
 ("Kinh tế–xã hội", "low_income", "Thu nhập thấp", "INDFMPIR (tỉ lệ thu nhập gia đình/ngưỡng nghèo)",
  "0 khi PIR ≥ 3; 1 khi PIR ≤ 1; tuyến tính ở giữa.",
  "Cân nhắc neo theo mốc thường dùng trong nghiên cứu NHANES: 1 khi PIR < 1,3; 0 khi PIR ≥ 3,5 [CẦN XÁC MINH mốc trong nguồn].",
  "Nghiên cứu thu nhập và triệu chứng trầm cảm NHANES 2005–2016 [R2]"),
 ("Kinh tế–xã hội", "low_education", "Học vấn thấp", "DMDEDUC2 (người ≥ 20 tuổi)",
  "Dưới lớp 9 = 1; lớp 9–11 = 1; tốt nghiệp THPT/GED = 0,5; một phần cao đẳng/đại học = 0; tốt nghiệp đại học trở lên = 0.",
  "Có thể giữ. Phương án khác: 'một phần cao đẳng' = 0,25 để thang có bốn bậc.", "—"),
 ("Kinh tế–xã hội", "no_partner", "Không có bạn đời", "DMDMARTL",
  "Góa, ly hôn, ly thân, chưa từng kết hôn = 1; đang kết hôn hoặc sống chung = 0.",
  "Cân nhắc tách 'chưa từng kết hôn' (ví dụ 0,5) khỏi nhóm mất bạn đời (góa, ly hôn, ly thân = 1), vì hai nhóm khác nhau về ý nghĩa lâm sàng.", "—"),
 ("Kinh tế–xã hội", "uninsured", "Không có bảo hiểm y tế", "HIQ011",
  "Không có bảo hiểm = 1; có = 0.", "Giữ (biến nhị phân theo định nghĩa).", "—"),
 ("Lối sống", "smoking", "Hút thuốc", "SMQ020, SMQ040",
  "Hút hằng ngày = 1; thỉnh thoảng = 0,5; đã bỏ = 0; chưa từng hút (< 100 điếu) = 0.",
  "Có thể giữ. Theo định nghĩa CDC, cả 'hằng ngày' và 'thỉnh thoảng' đều là người đang hút; mức 0,5 là lựa chọn phân bậc của tác giả. Cân nhắc 'đã bỏ' = 0,25 nếu muốn phản ánh tiền sử.",
  "Định nghĩa tình trạng hút thuốc của NHIS/CDC [R3]"),
 ("Lối sống", "heavy_drinking", "Uống rượu nhiều", "ALQ121/ALQ120Q, ALQ130",
  "Không uống trong năm = 0; với người có uống: 0 khi ≤ 2 ly/ngày uống, 1 khi ≥ 5 ly, tuyến tính ở giữa. Không tách theo giới.",
  "Cân nhắc tách theo giới theo NIAAA (uống nhiều: nam > 4 ly/ngày, nữ > 3 ly/ngày). Ví dụ: nam 0 tại 2, 1 tại 5; nữ 0 tại 1, 1 tại 4. Cần thêm biến giới RIAGENDR.",
  "NIAAA, Drinking Levels Defined [R4]"),
 ("Lối sống", "inactivity", "Ít vận động", "PAQ605, PAQ620, PAQ650, PAQ665",
  "Có vận động mạnh (công việc hoặc giải trí) = 0; chỉ vận động vừa = 0,5; không vận động = 1.",
  "Có thể giữ. Các câu PAQ chỉ hỏi có/không (≥ 10 phút liên tục), không đủ để tính 150 phút/tuần theo khuyến nghị; nên ghi là hạn chế.",
  "Physical Activity Guidelines for Americans, 2nd ed. [R5]"),
 ("Lối sống", "short_sleep", "Ngủ ngắn", "SLD012 / SLD010H",
  "0 khi ngủ ≥ 7 giờ; 1 khi ngủ ≤ 5 giờ; tuyến tính ở giữa.",
  "Mốc 7 giờ khớp khuyến nghị AASM/SRS. Cân nhắc mốc trên 6 giờ thay vì 5. Ngủ dài (> 9 giờ) cũng liên quan trầm cảm nhưng chưa có trong mô hình; có thể thêm khái niệm long_sleep.",
  "Watson et al. 2015 [R6]"),
 ("Lối sống", "obesity", "Béo phì", "BMXBMI (đo trực tiếp)",
  "0 khi BMI ≤ 25; 1 khi BMI ≥ 35; tuyến tính (μ = 0,5 tại BMI 30).",
  "Có thể giữ (μ = 0,5 trùng mốc béo phì WHO). Phương án khác: 0 tại 25, 1 tại 30, nếu muốn μ = 1 cho mọi người béo phì.",
  "WHO Technical Report Series 894 [R7]"),
 ("Sức khỏe", "poor_self_rated_health", "Sức khỏe tự đánh giá kém", "HSD010 (1 = rất tốt … 5 = kém)",
  "Tuyến tính: 1 → 0; 2 → 0,25; 3 → 0,5; 4 → 0,75; 5 → 1.",
  "Có thể giữ. Nhiều nghiên cứu gộp 'trung bình/kém' thành nhóm nguy cơ; nếu chuyên gia muốn, dùng: 1–2 = 0; 3 = 0,33; 4 = 0,67; 5 = 1.",
  "Idler & Benyamini 1997 [R8]"),
 ("Sức khỏe", "hypertension", "Tăng huyết áp", "BPQ020",
  "Từng được bác sĩ nói bị tăng huyết áp = 1; không = 0.", "Giữ (tự báo cáo chẩn đoán).", "—"),
 ("Sức khỏe", "diabetes", "Đái tháo đường", "DIQ010",
  "Có chẩn đoán = 1; 'giáp ranh' (borderline) = 0,5; không = 0.",
  "Có thể giữ. 'Giáp ranh' tương ứng tiền đái tháo đường; cân nhắc 0,3–0,5.",
  "ADA, Standards of Care – Diagnosis and Classification [R9]"),
 ("Sức khỏe", "cardiovascular", "Bệnh tim mạch", "MCQ160B, C, E, F",
  "Có ít nhất một trong: suy tim, bệnh mạch vành, nhồi máu cơ tim, đột quỵ = 1; không = 0.",
  "Giữ. Cân nhắc có nên tách đột quỵ (bệnh mạch máu não) ra không.", "—"),
 ("Sức khỏe", "cancer", "Ung thư", "MCQ220",
  "Từng được chẩn đoán ung thư bất kỳ = 1; không = 0.",
  "Cân nhắc loại trừ ung thư da không phải melanoma (dùng MCQ230 để xác định loại), vì nhóm này phổ biến và ít gánh nặng tâm lý.", "—"),
 ("Sức khỏe", "arthritis", "Viêm khớp", "MCQ160A",
  "Từng được bác sĩ nói bị viêm khớp = 1; không = 0.", "Giữ.", "—"),
]
for i, r in enumerate(R, 1):
    ws.append([i] + list(r) + ["", "", "", ""])
n = len(R) + 1
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.font = Font(name=F, size=11); c.alignment = wrap; c.border = B
    for c in row[6:8]:
        c.fill = SUG; c.font = Font(name=F, size=10, italic=True, color="404040")
    for c in row[8:12]:
        c.fill = FILL
widths = [5, 14, 22, 18, 20, 40, 46, 26, 14, 36, 24, 26]
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
dv = DataValidation(type="list", formula1='"Giữ,Sửa"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"I2:I{n}")
ws.freeze_panes = "D2"
ws.cell(n + 2, 1, "Đã duyệt:").font = Font(name=F, bold=True)
ws.cell(n + 2, 3, f'=COUNTA(I2:I{n})&" / {len(R)} dòng"').font = Font(name=F)
ws.cell(n + 3, 1, "Số dòng 'Sửa':").font = Font(name=F, bold=True)
ws.cell(n + 3, 3, f'=COUNTIF(I2:I{n},"Sửa")').font = Font(name=F)

# ---------------- Sheet 3: Câu hỏi lâm sàng ----------------
ws = wb.create_sheet("Cau hoi lam sang")
ws.append(["STT", "Kết quả từ dữ liệu (pipeline hiện tại, trước khi duyệt)", "Câu hỏi cho chuyên gia", "Ý kiến chuyên gia"])
style_header(ws, 1, 4)
Q = [
 ("Thuật toán Leiden chia 24 khái niệm thành 3 cộng đồng: (a) triệu chứng trầm cảm; (b) sức khỏe thực thể; (c) bất lợi kinh tế–xã hội. Cách chia giữ nguyên trên dữ liệu 2007–2012.",
  "Cách chia này có hợp lý về lâm sàng không? Có khái niệm nào nằm sai nhóm theo kinh nghiệm của cô?"),
 ("Khi bỏ trọng số (giữ mọi cạnh với trọng số 1), mệt mỏi và rối loạn giấc ngủ chuyển sang cộng đồng sức khỏe thực thể.",
  "Điều này có khớp với cách phân biệt triệu chứng cơ thể (somatic) và triệu chứng nhận thức–cảm xúc của trầm cảm không? Có tài liệu nào cô muốn dẫn?"),
 ("15/24 khái niệm có cạnh đi ra mạnh nhất trỏ về 'sức khỏe tự đánh giá kém'.",
  "Theo cô, đây là một nút trung gian có ý nghĩa lâm sàng, hay chủ yếu là do tỉ lệ cao của khái niệm này trong dân số?"),
 ("Bài báo không khẳng định quan hệ nhân quả và không đánh giá giá trị chẩn đoán; dữ liệu là cắt ngang.",
  "Cô đề xuất 1–2 câu cho phần Discussion về ý nghĩa thực hành (ví dụ cho sàng lọc ở tuyến cơ sở), trong giới hạn này."),
]
for i, (a, b) in enumerate(Q, 1):
    ws.append([i, a, b, ""])
for row in ws.iter_rows(min_row=2, max_row=len(Q) + 1):
    for c in row:
        c.font = Font(name=F, size=11); c.alignment = wrap; c.border = B
    row[3].fill = FILL
    ws.row_dimensions[row[0].row].height = 90
for col, w in zip("ABCD", [5, 55, 50, 60]):
    ws.column_dimensions[col].width = w
r0 = len(Q) + 3
for k, t in enumerate(["Người duyệt (họ tên, học vị):", "Đơn vị:", "Ngày duyệt:", "Chữ ký:"]):
    ws.cell(r0 + k, 2, t).font = Font(name=F, bold=True)
    ws.cell(r0 + k, 3).fill = FILL; ws.cell(r0 + k, 3).border = B
ws.cell(r0 + 1, 3, "Nguyen Tat Thanh University").font = Font(name=F)

# ---------------- Sheet 4: Nguồn ----------------
ws = wb.create_sheet("Nguon")
ws.append(["Mã", "Tài liệu", "DOI / URL", "Ghi chú"]); style_header(ws, 1, 4)
S = [
 ("R1", "Kroenke K, Spitzer RL, Williams JBW. The PHQ-9: validity of a brief depression severity measure. J Gen Intern Med. 2001;16(9):606–613.", "https://doi.org/10.1046/j.1525-1497.2001.016009606.x", ""),
 ("R2", "How Income and Income Inequality Drive Depressive Symptoms in U.S. Adults, Does Sex Matter: 2005–2016. Int J Environ Res Public Health. 2022.", "https://www.ncbi.nlm.nih.gov/pmc/articles/PMC9140340/", "[CẦN XÁC MINH] mốc PIR dùng trong bài"),
 ("R3", "CDC/NCHS. National Health Interview Survey – Adult tobacco use glossary (current, former, never smoker).", "https://www.cdc.gov/nchs/nhis/tobacco/tobacco_glossary.htm", "[CẦN XÁC MINH] URL"),
 ("R4", "National Institute on Alcohol Abuse and Alcoholism. Drinking Levels Defined.", "https://www.niaaa.nih.gov/alcohol-health/overview-alcohol-consumption/moderate-binge-drinking", "[CẦN XÁC MINH] URL"),
 ("R5", "Piercy KL, Troiano RP, Ballard RM, et al. The Physical Activity Guidelines for Americans. JAMA. 2018;320(19):2020–2028.", "https://doi.org/10.1001/jama.2018.14854", ""),
 ("R6", "Watson NF, Badr MS, Belenky G, et al. Recommended amount of sleep for a healthy adult: a joint consensus statement of the AASM and SRS. Sleep. 2015;38(6):843–844.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4434546/", ""),
 ("R7", "World Health Organization. Obesity: preventing and managing the global epidemic. WHO Technical Report Series 894. Geneva: WHO; 2000.", "https://iris.who.int/handle/10665/42330", "[CẦN XÁC MINH] URL"),
 ("R8", "Idler EL, Benyamini Y. Self-rated health and mortality: a review of twenty-seven community studies. J Health Soc Behav. 1997;38(1):21–37.", "https://pubmed.ncbi.nlm.nih.gov/9097506/", ""),
 ("R9", "American Diabetes Association. Diagnosis and Classification of Diabetes: Standards of Care in Diabetes (bản hiện hành).", "https://diabetesjournals.org/care", "[CẦN XÁC MINH] năm và trang"),
]
for s in S:
    ws.append(list(s))
for row in ws.iter_rows(min_row=2, max_row=len(S) + 1):
    for c in row:
        c.font = Font(name=F, size=11); c.alignment = wrap; c.border = B
for col, w in zip("ABCD", [6, 80, 55, 30]):
    ws.column_dimensions[col].width = w

wb.save("Phieu_duyet_ham_thanh_vien_VanAn.xlsx")
